from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Quality_Sweets_First_Month_Progress_Report_September_2026.docx"

NAVY = "24323D"
SAFFRON = "C9851C"
PALE_GOLD = "F8F0DE"
PALE_BLUE = "EDF3F6"
PALE_GREEN = "EAF3EC"
LIGHT = "D9D9D9"
MID = "6B7280"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=LIGHT, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')


def set_font(run, size=10, bold=False, color='000000', name='Aptos'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_run(p, text, size=10, bold=False, color='000000'):
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return r


def add_bullet(doc, text, level=0, size=9.4):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.18 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.04
    add_run(p, text, size)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, text, 13, True, '000000')
    return p


def add_label(cell, title, body, fill):
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    set_cell_margins(cell, 130, 145, 125, 145)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_run(p, title, 10, True, NAVY)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.03
    add_run(p, body, 8.6, False, '374151')


def add_timeline_row(table, date, title, body, fill):
    cells = table.add_row().cells
    set_col_width(cells[0], 1.12)
    set_col_width(cells[1], 5.75)
    for cell in cells:
        set_cell_border(cell)
        set_cell_margins(cell, 90, 115, 90, 115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cells[0], fill)
    p = cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, date, 8.6, True, NAVY)
    p = cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    add_run(p, title, 9.7, True, '000000')
    p = cells[1].add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    add_run(p, body, 8.65, False, '374151')


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.56)
sec.bottom_margin = Inches(0.54)
sec.left_margin = Inches(0.68)
sec.right_margin = Inches(0.68)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10)

# Header
header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_after = Pt(0)
add_run(hp, 'QUALITY SWEETS  |  FIRST MONTH UPDATE', 7.5, True, MID)

# Title and opening
p = doc.add_paragraph(style='Title')
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 0.98
add_run(p, 'First Month Progress Report', 23, True, '000000')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
add_run(p, 'September 2026', 10.5, True, SAFFRON)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
p.paragraph_format.line_spacing = 1.1
add_run(p, 'This first month focused on building a stronger online foundation for Quality Sweets. The work completed makes it easier for customers to find the business, understand the offer, and take the next step online. We will begin tracking results as the new foundation has time to take effect.', 10.2, False, '2F3A42')

section_heading(doc, 'What we put in place')
overview = doc.add_table(rows=2, cols=2)
overview.alignment = WD_TABLE_ALIGNMENT.CENTER
overview.autofit = False
for row in overview.rows:
    for c in row.cells:
        set_col_width(c, 3.44)
add_label(overview.cell(0, 0), 'Clearer website journey', 'A simpler path from the homepage to sweets, takeout, local pickup and catering.', PALE_GOLD)
add_label(overview.cell(0, 1), 'Cleaner product catalogue', 'Product names, page addresses and old links were cleaned up to reduce confusion.', PALE_BLUE)
add_label(overview.cell(1, 0), 'Stronger local presence', 'The Google Business Profile now reflects the store, takeout service and key offerings more clearly.', PALE_GREEN)
add_label(overview.cell(1, 1), 'Ready to measure progress', 'Tracking is in place so upcoming reports can focus on meaningful customer actions.', 'F3F4F6')

section_heading(doc, 'Work completed during the month')
timeline = doc.add_table(rows=0, cols=2)
timeline.alignment = WD_TABLE_ALIGNMENT.CENTER
timeline.autofit = False
add_timeline_row(timeline, 'Sep 10-15', 'Website and online ordering foundation', 'Improved the main website structure; created clearer pages for takeout, the Iselin store and catering; refreshed the homepage; and added helpful content to the main sweets and gift categories.', PALE_GOLD)
add_timeline_row(timeline, 'Sep 10-15', 'Cleaner catalogue and easier navigation', 'Updated category page addresses, cleaned eight pickup-only product pages, and made sure older links lead customers to the right current pages.', PALE_BLUE)
add_timeline_row(timeline, 'Sep 16-20', 'Local business information and menu', 'Aligned the store details used online, clarified that service is takeout and pickup only, and updated the Google Business Profile with products and menu information.', PALE_GREEN)
add_timeline_row(timeline, 'Sep 16-20', 'Measurement setup', 'Added tracking for purchases, catering inquiries, calls, emails, directions and menu visits. This creates a reliable starting point for future reports.', 'F3F4F6')
add_timeline_row(timeline, 'Sep 29', 'More product visibility on Google', 'Connected the online catalogue to Google Shopping free listings so eligible sweets, snacks and gift boxes can be discovered in Google.', PALE_GOLD)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(0)
add_run(p, 'This report intentionally focuses on the work completed, not early performance results. The next reporting period will establish the first comparable baseline.', 8.8, False, MID)

# Second page
doc.add_page_break()
section_heading(doc, 'What happens next')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.line_spacing = 1.08
add_run(p, 'With the core pages, local listing and tracking in place, the next phase is about helping more people discover Quality Sweets and measuring what they do once they arrive.', 10, False, '2F3A42')

next_table = doc.add_table(rows=1, cols=3)
next_table.alignment = WD_TABLE_ALIGNMENT.CENTER
next_table.autofit = False
headers = ['Next priority', 'What it means for Quality Sweets', 'Status']
widths = [1.7, 4.25, 0.93]
for i, h in enumerate(headers):
    c = next_table.cell(0, i)
    set_col_width(c, widths[i])
    set_cell_shading(c, NAVY)
    set_cell_border(c, NAVY)
    set_cell_margins(c, 100, 105, 100, 105)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 2 else WD_ALIGN_PARAGRAPH.LEFT
    add_run(p, h, 8.5, True, WHITE)

next_rows = [
    ('Help Google understand the website', 'Add the behind-the-scenes website information that helps Google recognize the business, products and page relationships.', 'Next'),
    ('Build local trust', 'Add real photos, regular Google updates and a simple review request process for store customers.', 'Next'),
    ('Keep business listings consistent', 'Check and correct key business directories so customers find the same address, phone number and website everywhere.', 'Next'),
    ('Watch customer actions', 'Review search visibility, website visits, calls, directions, menu views, online purchases and catering inquiries.', 'Ongoing'),
    ('Prepare seasonal opportunities', 'Create useful pages and offers for major celebrations when products and timing are confirmed.', 'Planned'),
]
for index, (priority, meaning, status) in enumerate(next_rows):
    cells = next_table.add_row().cells
    fills = ['FFFFFF', 'FFFFFF', PALE_GOLD if status == 'Next' else PALE_GREEN if status == 'Ongoing' else PALE_BLUE]
    for i, cell in enumerate(cells):
        set_col_width(cell, widths[i])
        set_cell_border(cell)
        set_cell_shading(cell, fills[i])
        set_cell_margins(cell, 95, 105, 95, 105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cells[0].paragraphs[0]
    p.paragraph_format.line_spacing = 1.02
    add_run(p, priority, 8.85, True, NAVY)
    p = cells[1].paragraphs[0]
    p.paragraph_format.line_spacing = 1.02
    add_run(p, meaning, 8.55, False, '374151')
    p = cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, status, 8.25, True, NAVY)

section_heading(doc, 'What future reports will show')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
add_run(p, 'Starting next month, we will use the new tracking to report on a small set of clear, useful indicators:', 9.8, False, '2F3A42')
for t in [
    'How easily customers find Quality Sweets on Google and Google Maps.',
    'Visits to the website and key pages such as takeout, the Iselin location and catering.',
    'Customer actions: calls, direction requests, menu views, online purchases and catering inquiries.',
    'The work completed during the period and the next priorities.'
]:
    add_bullet(doc, t, size=9.1)

section_heading(doc, 'The focus from here')
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.08
add_run(p, 'The groundwork is now in place. Our next step is to build visibility steadily, keep the online information current, and use real customer activity to guide each monthly update.', 10, False, '2F3A42')

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(4)
add_run(fp, 'Quality Sweets  |  First Month Progress Report  |  September 2026', 7.4, False, MID)

doc.save(OUT)
print(OUT)
