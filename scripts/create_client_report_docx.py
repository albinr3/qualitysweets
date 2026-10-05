from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Paths
BASE_DIR = Path(r"c:\Users\Albin Rodriguez\Documents\QualitySweets")
BASELINE = BASE_DIR / "baseline"
HEATMAP = Path(r"C:\Users\Albin Rodriguez\Downloads\Mapa de calor SEO en Nueva Jersey.png")

OUTPUT_DOCX_1 = BASE_DIR / "outputs" / "docx" / "Quality_Sweets_First_Month_Progress_Report_September_2026_native.docx"
OUTPUT_DOCX_FINAL = BASE_DIR / "outputs" / "docx" / "Quality_Sweets_First_Month_Progress_Report_September_2026.docx"
OUTPUT_SAME_DIR = BASE_DIR / "outputs" / "pdf" / "Quality_Sweets_First_Month_Progress_Report_September_2026.docx"

# Color constants
NAVY_HEX = "24323D"
SAFFRON_HEX = "C9851C"
GOLD_HEX = "F8F0DE"
BLUE_HEX = "EDF3F6"
GREEN_HEX = "EAF3EC"
GREY_HEX = "F3F4F6"
MID_HEX = "6B7280"
DARK_HEX = "2F3A42"
BORDER_HEX = "D9D9D9"
WHITE_HEX = "FFFFFF"

def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill_hex)

def set_cell_border(cell, color=BORDER_HEX, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')

def set_font(run, size=10, bold=False, color='000000', name='Helvetica'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_run(p, text, size=10, bold=False, color='000000', name='Helvetica'):
    r = p.add_run(text)
    set_font(r, size, bold, color, name)
    return r

def add_bullet(doc, text, size=9.3, color=DARK_HEX):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_run(p, text, size, False, color)
    return p

def create_document():
    doc = Document()
    
    # Page setup - Letter, 0.65 in left/right margins, 0.55 in top, 0.58 in bottom
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.58)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    sec.different_first_page_header_footer = False

    # Header
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    add_run(hp, "QUALITY SWEETS  |  FIRST MONTH UPDATE", 7.4, True, MID_HEX)

    # Footer
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)
    add_run(fp, "Quality Sweets  |  First Month Progress Report", 7.2, False, MID_HEX)

    # PAGE 1:
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    add_run(p_title, "First Month Progress Report", 23, True, "000000")

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.paragraph_format.line_spacing = 1.15
    add_run(p_intro, 
            "This report records the work completed during the first month: website updates, "
            "a clearer product catalogue, local business information, and tracking for customer actions.", 
            10.2, False, DARK_HEX)

    # Heading 1
    p_h2_1 = doc.add_paragraph()
    p_h2_1.paragraph_format.space_before = Pt(4)
    p_h2_1.paragraph_format.space_after = Pt(4)
    add_run(p_h2_1, "What we put in place", 13, True, "000000")

    # 2x2 cards table
    card_table = doc.add_table(rows=2, cols=2)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_table.autofit = False

    cards_data = [
        ("Clearer website journey", "A simpler path from the homepage to sweets, takeout, local pickup and catering.", GOLD_HEX),
        ("Cleaner product catalogue", "Product names, page addresses and old links were cleaned up to reduce confusion.", BLUE_HEX),
        ("Stronger local presence", "The Google Business Profile now reflects the store, takeout service and key offerings more clearly.", GREEN_HEX),
        ("Ready to measure activity", "Tracking is in place for customer actions such as calls, directions, menu visits and inquiries.", GREY_HEX),
    ]

    for idx, (ctitle, cbody, fill) in enumerate(cards_data):
        r = idx // 2
        c = idx % 2
        cell = card_table.cell(r, c)
        set_col_width(cell, 3.45)
        set_cell_shading(cell, fill)
        set_cell_border(cell, BORDER_HEX, '6')
        set_cell_margins(cell, top=120, start=135, bottom=120, end=135)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.05
        add_run(cp, ctitle, 9.8, True, NAVY_HEX)
        
        cp2 = cell.add_paragraph()
        cp2.paragraph_format.space_after = Pt(0)
        cp2.paragraph_format.line_spacing = 1.08
        add_run(cp2, cbody, 8.55, False, "374151")

    # Heading 2
    p_h2_2 = doc.add_paragraph()
    p_h2_2.paragraph_format.space_before = Pt(10)
    p_h2_2.paragraph_format.space_after = Pt(4)
    add_run(p_h2_2, "Work completed", 13, True, "000000")

    # Work table (5 rows, 2 columns)
    work = [
        ("Website and online ordering", "Clearer pages for customers", "Improved the website structure; created clearer pages for takeout, the Iselin store and catering; refreshed the homepage; and added useful information to the main sweets and gift categories.", GOLD_HEX),
        ("Product catalogue", "Cleaner catalogue and easier navigation", "Updated category page addresses, cleaned eight pickup-only product pages, and made sure older links lead customers to the current pages.", BLUE_HEX),
        ("Local business profile", "Store details, products and menu", "Aligned the store details used online, clarified takeout and pickup service, and updated the Google Business Profile with products and menu information.", GREEN_HEX),
        ("Measurement", "Customer-action tracking", "Added tracking for purchases, catering inquiries, calls, emails, directions and menu visits.", GREY_HEX),
        ("Google Shopping", "More product visibility on Google", "Connected eligible sweets, snacks and gift boxes to Google Shopping free listings.", GOLD_HEX),
    ]

    work_table = doc.add_table(rows=len(work), cols=2)
    work_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    work_table.autofit = False

    for i, (area_text, title_text, body_text, fill) in enumerate(work):
        c0 = work_table.cell(i, 0)
        c1 = work_table.cell(i, 1)

        set_col_width(c0, 1.8)
        set_col_width(c1, 5.1)

        set_cell_shading(c0, fill)
        set_cell_border(c0, BORDER_HEX, '6')
        set_cell_margins(c0, top=100, start=110, bottom=100, end=110)
        c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        set_cell_shading(c1, WHITE_HEX)
        set_cell_border(c1, BORDER_HEX, '6')
        set_cell_margins(c1, top=80, start=110, bottom=80, end=110)
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        # Left cell content
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.05
        add_run(p0, area_text, 9.0, True, NAVY_HEX)

        # Right cell content
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1.5)
        p1.paragraph_format.line_spacing = 1.05
        add_run(p1, title_text, 9.5, True, "000000")

        p2 = c1.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.08
        add_run(p2, body_text, 8.5, False, "374151")

    # Footnote 1
    p_foot1 = doc.add_paragraph()
    p_foot1.paragraph_format.space_before = Pt(8)
    p_foot1.paragraph_format.space_after = Pt(0)
    add_run(p_foot1, "The next report is expected around November 1, 2026. It will compare the first full cycle of visibility and customer-action data.", 8.7, False, MID_HEX)

    # PAGE 2:
    doc.add_page_break()

    p_h2_p2_1 = doc.add_paragraph()
    p_h2_p2_1.paragraph_format.space_before = Pt(6)
    p_h2_p2_1.paragraph_format.space_after = Pt(4)
    add_run(p_h2_p2_1, "Google Maps starting point", 13, True, "000000")

    p_p2_1 = doc.add_paragraph()
    p_p2_1.paragraph_format.space_after = Pt(10)
    p_p2_1.paragraph_format.line_spacing = 1.15
    add_run(p_p2_1, "The following grid scans record current Google Maps positions for the tracked searches. They provide a starting reference for later reports.", 10.1, False, DARK_HEX)

    p_h2_p2_2 = doc.add_paragraph()
    p_h2_p2_2.paragraph_format.space_before = Pt(6)
    p_h2_p2_2.paragraph_format.space_after = Pt(4)
    add_run(p_h2_p2_2, "Why this report does not include commercial changes", 13, True, "000000")

    p_p2_2 = doc.add_paragraph()
    p_p2_2.paragraph_format.space_after = Pt(10)
    p_p2_2.paragraph_format.line_spacing = 1.15
    add_run(p_p2_2, "Sales, revenue and lead changes are not included because this period was used to prepare the website, Google Business Profile and tracking. The updated pages and listing need a full reporting cycle before customer activity can be compared. The grid scans show local visibility, not sales results.", 10.1, False, DARK_HEX)

    p_h2_p2_3 = doc.add_paragraph()
    p_h2_p2_3.paragraph_format.space_before = Pt(6)
    p_h2_p2_3.paragraph_format.space_after = Pt(4)
    add_run(p_h2_p2_3, "What happens next", 13, True, "000000")

    for item in [
        "Add the website information Google needs to better understand the business, products and pages.",
        "Add real photos, regular Google updates and a simple review request process.",
        "Review customer actions and use the data to guide the next priorities."
    ]:
        add_bullet(doc, item, size=9.3, color=DARK_HEX)

    p_foot2 = doc.add_paragraph()
    p_foot2.paragraph_format.space_before = Pt(12)
    p_foot2.paragraph_format.space_after = Pt(0)
    add_run(p_foot2, "The grid scans and New Jersey heat map follow. Each image is included in its original form.", 8.7, False, MID_HEX)

    # PAGES 3 to 10:
    screens = [
        ("Indian sweets shop near me", BASELINE / "Captura de pantalla 2026-10-03 163017.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Indian sweets near me", BASELINE / "Captura de pantalla 2026-10-03 163116.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Samosas near me", BASELINE / "Captura de pantalla 2026-10-03 163149.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Samosa chaat near me", BASELINE / "Captura de pantalla 2026-10-03 163212.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Chaat near me", BASELINE / "Captura de pantalla 2026-10-03 163257.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Bengali sweets", BASELINE / "Captura de pantalla 2026-10-03 163320.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("Pani puri near me", BASELINE / "Captura de pantalla 2026-10-03 163354.png", "Google Maps grid scan", "Original Google Maps grid scan for this tracked search."),
        ("New Jersey SEO Heat Map", HEATMAP, "New Jersey local search scan", "Original New Jersey SEO heat map.")
    ]

    for title_text, img_path, subtitle_text, caption_text in screens:
        doc.add_page_break()

        # Image page title
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(4)
        p_t.paragraph_format.space_after = Pt(2)
        add_run(p_t, title_text, 18, True, "000000")

        # Subtitle
        p_s = doc.add_paragraph()
        p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_s.paragraph_format.space_after = Pt(10)
        add_run(p_s, subtitle_text, 9, False, MID_HEX)

        # Image
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        # Width: 6.8 inches fits letter width (8.5 - 2*0.65 = 7.2 max width)
        p_img.add_run().add_picture(str(img_path), width=Inches(6.8))

        # Caption
        p_c = doc.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(0)
        add_run(p_c, caption_text, 8.2, False, MID_HEX)

    # Save to both target destinations
    OUTPUT_DOCX_FINAL.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX_FINAL))
    print(f"Saved: {OUTPUT_DOCX_FINAL}")

    OUTPUT_SAME_DIR.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_SAME_DIR))
    print(f"Saved: {OUTPUT_SAME_DIR}")

if __name__ == "__main__":
    create_document()
